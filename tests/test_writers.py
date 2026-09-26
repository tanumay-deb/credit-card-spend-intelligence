import csv
import sys
from datetime import date
from decimal import Decimal
from pathlib import Path

import pytest

from creditcard import writers
from creditcard.models import IngestLogRow, StatementSummary, Transaction
from creditcard.writers import (
    OutputLocked,
    is_locked,
    write_excel_workbook,
    write_ingest_log,
    write_outputs,
    write_review_queue,
    write_statements,
    write_transactions,
)


def _txn(txn_id="a1", posting_date=date(2026, 8, 4)):
    return Transaction(
        txn_id=txn_id, card_id="c1", statement_id="s1", txn_date=date(2026, 8, 3),
        posting_date=posting_date, merchant_raw="SWIGGY*ORDER",
        merchant_clean="SWIGGY", amount=Decimal("482.50"), direction="debit",
        txn_type="purchase", category_group="Food", subcategory="Swiggy",
        category_source="rule", is_forex=False, forex_ccy=None,
        forex_amount=None, source_file="f.pdf",
    )


def _read(path):
    with open(path, newline="", encoding="utf-8-sig") as fh:
        return list(csv.DictReader(fh))


def test_write_transactions_produces_readable_csv(tmp_path):
    out = tmp_path / "transactions.csv"
    write_transactions([_txn()], out)
    rows = _read(out)
    assert rows[0]["txn_id"] == "a1"
    assert rows[0]["amount"] == "482.50"
    assert rows[0]["txn_date"] == "2026-08-03"
    assert rows[0]["is_forex"] == "False"


def test_empty_input_still_writes_header(tmp_path):
    """Power BI must find a schema even before any statement is imported."""
    out = tmp_path / "transactions.csv"
    write_transactions([], out)
    text = out.read_text(encoding="utf-8-sig")
    assert text.startswith("txn_id,card_id")


def test_none_posting_date_writes_empty_not_the_word_none(tmp_path):
    out = tmp_path / "t.csv"
    write_transactions([_txn(posting_date=None)], out)
    assert _read(out)[0]["posting_date"] == ""


def test_column_order_matches_the_dataclass(tmp_path):
    """Power BI binds by position as well as name; order must be stable."""
    from dataclasses import fields
    out = tmp_path / "t.csv"
    write_transactions([_txn()], out)
    header = out.read_text(encoding="utf-8-sig").splitlines()[0].split(",")
    assert header == [f.name for f in fields(Transaction)]


def test_creates_missing_parent_directory(tmp_path):
    out = tmp_path / "nested" / "deeper" / "transactions.csv"
    write_transactions([_txn()], out)
    assert out.exists()


def test_utf8_merchant_name_survives_round_trip(tmp_path):
    out = tmp_path / "t.csv"
    txn = _txn()
    object.__setattr__(txn, "merchant_clean", "CAFÉ MOCHÁ")
    write_transactions([txn], out)
    assert _read(out)[0]["merchant_clean"] == "CAFÉ MOCHÁ"


def test_write_statements(tmp_path):
    stmt = StatementSummary(
        statement_id="s1", card_id="c1", statement_date=date(2026, 8, 18),
        due_date=date(2026, 9, 7), period_start=date(2026, 7, 19),
        period_end=date(2026, 8, 18), previous_balance=Decimal("0"),
        payments=Decimal("0"), purchases=Decimal("1000"), total_due=Decimal("1000"),
        min_due=Decimal("50"), finance_charges=Decimal("0"), late_fee=Decimal("0"),
        credit_limit=Decimal("100000"), available_limit=Decimal("99000"),
        source_file="f.pdf",
    )
    out = tmp_path / "statements.csv"
    write_statements([stmt], out)
    assert _read(out)[0]["total_due"] == "1000"


def test_write_review_queue(tmp_path):
    out = tmp_path / "review_queue.csv"
    write_review_queue(
        [{"merchant_clean": "NEW", "times_seen": 2, "total_amount": Decimal("300"),
          "example_txn_date": date(2026, 8, 1), "example_card": "c1"}],
        out,
    )
    assert _read(out)[0]["times_seen"] == "2"


def test_write_ingest_log(tmp_path):
    row = IngestLogRow(
        source_file="f.pdf", file_hash="abc123", parser="hdfc",
        run_timestamp="2026-09-08T15:00:00", rows_found=12,
        statement_purchases=Decimal("1000"), sum_check_delta=Decimal("0"),
        status="ok", error="",
    )
    out = tmp_path / "ingest_log.csv"
    write_ingest_log([row], out)
    parsed = _read(out)[0]
    assert parsed["status"] == "ok"
    assert parsed["sum_check_delta"] == "0"


def test_ingest_log_none_deltas_write_empty(tmp_path):
    """A failed parse has no delta; it must not write the word None."""
    row = IngestLogRow(
        source_file="f.pdf", file_hash="abc123", parser="hdfc",
        run_timestamp="2026-09-08T15:00:00", rows_found=0,
        statement_purchases=None, sum_check_delta=None,
        status="failed", error="wrong password",
    )
    out = tmp_path / "ingest_log.csv"
    write_ingest_log([row], out)
    parsed = _read(out)[0]
    assert parsed["sum_check_delta"] == ""
    assert parsed["statement_purchases"] == ""


def test_write_excel_workbook(tmp_path):
    import openpyxl

    out = tmp_path / "credit_card_data.xlsx"
    res = write_excel_workbook(
        transactions=[_txn()],
        statements=[],
        emi_plans=[],
        payments=[],
        review=[{"merchant_clean": "SWIGGY", "times_seen": 1, "total_amount": Decimal("482.50")}],
        ingest_log=[],
        path=out,
    )
    assert res == out
    assert out.exists()

    wb = openpyxl.load_workbook(out)
    assert "Transactions" in wb.sheetnames
    assert "Statements" in wb.sheetnames
    assert "EMI_Plans" in wb.sheetnames
    assert "Payments" in wb.sheetnames
    assert "Review_Queue" in wb.sheetnames
    assert "Ingest_Log" in wb.sheetnames

    # Check transactions row
    ws = wb["Transactions"]
    assert ws.cell(row=2, column=1).value == "a1"
    assert ws.cell(row=2, column=7).value == "SWIGGY"
    assert ws.cell(row=2, column=8).value == 482.50


def test_transactions_csv_carries_spend_amount(tmp_path):
    """The dashboard sums this column, so it has to reach the CSV."""
    from dataclasses import replace
    out = tmp_path / "transactions.csv"
    write_transactions([replace(_txn(), spend_amount=Decimal("482.50"))], out)
    assert out.read_text(encoding="utf-8-sig").splitlines()[0].split(",")[-1] == "spend_amount"
    assert _read(out)[0]["spend_amount"] == "482.50"



# --- write_outputs: stage everything, then replace, never half ---------------


CSV_NAMES = [
    "transactions.csv", "statements.csv", "review_queue.csv",
    "ingest_log.csv", "emi_plans.csv", "payments.csv",
]


def _outputs(out, **overrides):
    kwargs = dict(transactions=[_txn()], statements=[], review=[], ingest_log=[],
                  emi_plans=[], payments=[])
    kwargs.update(overrides)
    return write_outputs(out, **kwargs)


def _snapshot(folder):
    return {p.name: p.read_bytes() for p in folder.iterdir()}


def test_write_outputs_writes_every_file_and_leaves_no_temporaries(tmp_path):
    result = _outputs(tmp_path)
    for name in CSV_NAMES + ["credit_card_data.xlsx"]:
        assert (tmp_path / name).exists(), name
    assert result == {"excel_file": tmp_path / "credit_card_data.xlsx",
                      "workbook_skipped": False}
    assert not list(tmp_path.glob("*.tmp.*"))


def test_a_locked_csv_replaces_nothing(tmp_path, monkeypatch):
    """Excel holding one CSV used to crash the run halfway, leaving this run's
    transactions next to last run's statements. Now nothing moves."""
    _outputs(tmp_path)
    before = _snapshot(tmp_path)
    locked = tmp_path / "statements.csv"
    monkeypatch.setattr(writers, "is_locked", lambda p: Path(p) == locked)
    with pytest.raises(OutputLocked) as info:
        _outputs(tmp_path, transactions=[_txn("b2")])
    assert info.value.path == locked
    assert _snapshot(tmp_path) == before


def test_a_locked_workbook_still_updates_the_csvs(tmp_path, monkeypatch):
    """The workbook is a convenience copy; Power BI reads the CSVs."""
    _outputs(tmp_path)
    workbook = tmp_path / "credit_card_data.xlsx"
    old_workbook = workbook.read_bytes()
    monkeypatch.setattr(writers, "is_locked", lambda p: Path(p) == workbook)
    result = _outputs(tmp_path, transactions=[_txn("b2")])
    assert result == {"excel_file": None, "workbook_skipped": True}
    assert _read(tmp_path / "transactions.csv")[0]["txn_id"] == "b2"
    assert workbook.read_bytes() == old_workbook
    assert not list(tmp_path.glob("*.tmp.*"))


def test_a_missing_file_is_not_locked_and_is_not_created(tmp_path):
    target = tmp_path / "absent.csv"
    assert not is_locked(target)
    assert not target.exists()


@pytest.mark.skipif(sys.platform != "win32", reason="Windows share modes")
def test_is_locked_detects_a_file_held_open_the_way_excel_holds_it(tmp_path):
    """Excel opens a file letting others read it but not write it."""
    import ctypes
    from ctypes import wintypes

    target = tmp_path / "held.csv"
    target.write_text("x")
    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    kernel32.CreateFileW.argtypes = [
        wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, wintypes.LPVOID,
        wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE,
    ]
    kernel32.CreateFileW.restype = wintypes.HANDLE
    generic_read, file_share_read, open_existing = 0x80000000, 0x1, 3
    handle = kernel32.CreateFileW(str(target), generic_read, file_share_read,
                                  None, open_existing, 0, None)
    assert handle not in (None, ctypes.c_void_p(-1).value)
    try:
        assert is_locked(target)
    finally:
        kernel32.CloseHandle(handle)
    assert not is_locked(target)
