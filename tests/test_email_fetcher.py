"""Unit tests for creditcard.email_fetcher."""
from __future__ import annotations

from decimal import Decimal
from email.mime.application import MIMEApplication
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from pathlib import Path
from unittest.mock import MagicMock

import pytest

from creditcard.email_fetcher import (
    DownloadResult,
    EmailConfig,
    EmailRule,
    build_search_query,
    decode_header_str,
    export_fetch_results,
    fetch_and_process_statements,
    load_email_config,
    load_email_rules,
    match_card_for_email,
    parse_email_message,
)
from creditcard.models import Card


def test_load_email_config(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    env_file = tmp_path / "secrets.env"
    env_file.write_text(
        "EMAIL_IMAP_SERVER=imap.example.com\n"
        "EMAIL_IMAP_PORT=993\n"
        "EMAIL_USERNAME=test@example.com\n"
        "EMAIL_PASSWORD=secretapppassword\n"
        "EMAIL_FOLDER=INBOX\n",
        encoding="utf-8-sig",
    )
    # Clear any existing env vars
    monkeypatch.delenv("EMAIL_USERNAME", raising=False)
    monkeypatch.delenv("EMAIL_PASSWORD", raising=False)
    monkeypatch.delenv("EMAIL_IMAP_SERVER", raising=False)

    cfg = load_email_config(env_file)
    assert cfg.host == "imap.example.com"
    assert cfg.port == 993
    assert cfg.username == "test@example.com"
    assert cfg.password == "secretapppassword"
    assert cfg.folder == "INBOX"


def test_load_email_config_missing_raises(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    env_file = tmp_path / "secrets.env"
    env_file.write_text("", encoding="utf-8-sig")
    monkeypatch.delenv("EMAIL_USERNAME", raising=False)
    monkeypatch.delenv("EMAIL_PASSWORD", raising=False)

    with pytest.raises(ValueError, match="EMAIL_USERNAME"):
        load_email_config(env_file)


def test_load_email_rules(tmp_path: Path):
    rules_file = tmp_path / "email_rules.csv"
    rules_file.write_text(
        "card_id,sender_contains,subject_contains,body_contains,filename_pattern\n"
        "hdfc_regalia,hdfcbank,Regalia,,*.pdf\n"
        "sbi_elite,sbicard,ELITE,,*.pdf\n",
        encoding="utf-8-sig",
    )
    rules = load_email_rules(rules_file)
    assert len(rules) == 2
    assert rules[0].card_id == "hdfc_regalia"
    assert rules[0].sender_contains == "hdfcbank"
    assert rules[0].subject_contains == "regalia"
    assert rules[1].card_id == "sbi_elite"


def test_decode_header_str():
    assert decode_header_str("Simple Subject") == "Simple Subject"
    # RFC 2047 encoded
    assert decode_header_str("=?utf-8?b?UmVnYWxpYSBIREZD?=") == "Regalia HDFC"
    assert decode_header_str(None) == ""


def test_match_card_for_email():
    rules = [
        EmailRule("hdfc_regalia", sender_contains="hdfc", subject_contains="regalia"),
        EmailRule("hdfc_millennia", sender_contains="hdfc", subject_contains="millennia"),
        EmailRule("sbi_bpcl", sender_contains="sbi", subject_contains="bpcl"),
    ]
    cards = {
        "hdfc_regalia": Card("hdfc_regalia", "HDFC", "Regalia", "1234", "hdfc", "P1", Decimal("0"), 1),
        "hdfc_millennia": Card("hdfc_millennia", "HDFC", "Millennia", "5678", "hdfc", "P2", Decimal("0"), 1),
        "sbi_elite": Card("sbi_elite", "SBI", "ELITE", "9012", "sbi", "P3", Decimal("0"), 1),
    }

    # 1. Matches rule
    assert match_card_for_email(
        "alerts@hdfcbank.net", "Your Regalia Credit Card Statement", "", "stmt.pdf", rules, cards
    ) == "hdfc_regalia"

    # 2. Matches via last4
    assert match_card_for_email(
        "alerts@sbicard.com", "Statement for card ending 9012", "", "stmt.pdf", rules, cards
    ) == "sbi_elite"

    # 3. Matches via card name
    assert match_card_for_email(
        "statement@sbicard.com", "SBI Card ELITE Statement", "", "stmt.pdf", rules, cards
    ) == "sbi_elite"

    # 4. Unmatched
    assert match_card_for_email(
        "newsletter@random.com", "Special Offers", "", "offers.pdf", rules, cards
    ) is None


def test_parse_email_message():
    msg = MIMEMultipart()
    msg["Subject"] = "Your HDFC Statement"
    msg["From"] = "statements@hdfcbank.net"
    msg["Date"] = "Wed, 10 Sep 2026 10:00:00 +0000"

    body = MIMEText("Attached is your monthly credit card statement.", "plain")
    msg.attach(body)

    pdf_attachment = MIMEApplication(b"%PDF-1.4 mock pdf content", _subtype="pdf")
    pdf_attachment.add_header("Content-Disposition", "attachment", filename="stmt_sep2026.pdf")
    msg.attach(pdf_attachment)

    raw_bytes = msg.as_bytes()
    date_iso, sender, subject, body_text, attachments = parse_email_message(raw_bytes)

    assert date_iso == "2026-09-10"
    assert "hdfcbank.net" in sender
    assert subject == "Your HDFC Statement"
    assert "Attached is your monthly" in body_text
    assert len(attachments) == 1
    fn, payload = attachments[0]
    assert fn == "stmt_sep2026.pdf"
    assert payload == b"%PDF-1.4 mock pdf content"


def test_build_search_query():
    assert "SINCE" in build_search_query(days_back=30)
    assert build_search_query(days_back=None, unread_only=False) == "ALL"
    assert "UNSEEN" in build_search_query(days_back=None, unread_only=True)


def test_find_candidate_messages_with_sender_and_subject():
    from creditcard.email_fetcher import find_candidate_messages

    mock_client = MagicMock()
    mock_client.search.return_value = ("OK", [b"42 99"])

    # Test explicit sender and subject filters
    ids = find_candidate_messages(
        client=mock_client,
        rules=[],
        cards={},
        days_back=30,
        sender_filter="hdfcbank",
        subject_filter="statement",
    )
    assert ids == ["42", "99"]
    call_query = mock_client.search.call_args[0][1]
    assert 'FROM "hdfcbank"' in call_query
    assert 'SUBJECT "statement"' in call_query
    assert "SINCE" in call_query


def test_fetch_and_process_statements_mock(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    cfg_dir = tmp_path / "config"
    cfg_dir.mkdir()
    (cfg_dir / "secrets.env").write_text(
        "EMAIL_USERNAME=user@test.com\nEMAIL_PASSWORD=pass\n", encoding="utf-8-sig"
    )
    (cfg_dir / "cards.csv").write_text(
        "card_id,issuer,card_name,last4,parser,password_env,credit_limit,statement_day,limit_group\n"
        "hdfc_regalia,HDFC,Regalia,0000,hdfc,P1,0,1,hdfc_regalia\n",
        encoding="utf-8-sig",
    )
    (cfg_dir / "email_rules.csv").write_text(
        "card_id,sender_contains,subject_contains,body_contains,filename_pattern\n"
        "hdfc_regalia,hdfc,regalia,,*.pdf\n",
        encoding="utf-8-sig",
    )

    # Build mock email with PDF
    msg = MIMEMultipart()
    msg["Subject"] = "Your Regalia HDFC Statement"
    msg["From"] = "statements@hdfcbank.net"
    msg["Date"] = "Wed, 10 Sep 2026 12:00:00 +0000"
    pdf = MIMEApplication(b"%PDF-1.4 statement content here", _subtype="pdf")
    pdf.add_header("Content-Disposition", "attachment", filename="hdfc_regalia_202609.pdf")
    msg.attach(pdf)

    # Mock IMAP client
    mock_client = MagicMock()
    mock_client.select.return_value = ("OK", [b"1"])
    mock_client.search.return_value = ("OK", [b"101"])
    mock_client.fetch.return_value = ("OK", [(b"101 (RFC822 {1234}", msg.as_bytes())])

    # Run fetch
    results = fetch_and_process_statements(
        root_dir=tmp_path,
        days_back=30,
        dry_run=False,
        client=mock_client,
    )

    assert len(results) == 1
    res = results[0]
    assert res.status == "downloaded"
    assert res.card_id == "hdfc_regalia"

    saved_file = tmp_path / "statements" / "hdfc_regalia" / "hdfc_regalia_202609.pdf"
    assert saved_file.exists()
    assert saved_file.read_bytes() == b"%PDF-1.4 statement content here"

    # Second run should skip due to existing hash
    results2 = fetch_and_process_statements(
        root_dir=tmp_path,
        days_back=30,
        dry_run=False,
        client=mock_client,
    )
    assert len(results2) == 1
    assert results2[0].status == "skipped_hash"


def test_parse_folders():
    from creditcard.email_fetcher import parse_folders

    assert parse_folders("") == ("INBOX",)
    assert parse_folders("INBOX") == ("INBOX",)
    assert parse_folders("INBOX, Bank Statements") == ("INBOX", "Bank Statements")
    assert parse_folders('INBOX, "Bank Statements", "[Gmail]/All Mail"') == (
        "INBOX",
        "Bank Statements",
        "[Gmail]/All Mail",
    )


def test_list_imap_folders():
    from creditcard.email_fetcher import list_imap_folders

    mock_client = MagicMock()
    mock_client.list.return_value = (
        "OK",
        [
            b'(\\HasNoChildren) "/" "INBOX"',
            b'(\\HasNoChildren) "/" "Bank Statements"',
            b'(\\HasNoChildren) "/" "[Gmail]/All Mail"',
        ],
    )
    folders = list_imap_folders(mock_client)
    assert folders == ["Bank Statements", "INBOX", "[Gmail]/All Mail"]


def test_multi_folder_fetch(tmp_path: Path):
    from creditcard.email_fetcher import fetch_and_process_statements

    cfg_dir = tmp_path / "config"
    cfg_dir.mkdir()
    (cfg_dir / "secrets.env").write_text(
        "EMAIL_USERNAME=user@test.com\nEMAIL_PASSWORD=pass\nEMAIL_FOLDERS=INBOX, \"Bank Statements\"\n",
        encoding="utf-8-sig",
    )
    (cfg_dir / "cards.csv").write_text(
        "card_id,issuer,card_name,last4,parser,password_env,credit_limit,statement_day,limit_group\n"
        "hdfc_regalia,HDFC,Regalia,0000,hdfc,P1,0,1,hdfc_regalia\n",
        encoding="utf-8-sig",
    )
    (cfg_dir / "email_rules.csv").write_text(
        "card_id,sender_contains,subject_contains,body_contains,filename_pattern\n"
        "hdfc_regalia,hdfc,regalia,,*.pdf\n",
        encoding="utf-8-sig",
    )

    msg = MIMEMultipart()
    msg["Subject"] = "Your Regalia HDFC Statement"
    msg["From"] = "statements@hdfcbank.net"
    msg["Date"] = "Wed, 10 Sep 2026 12:00:00 +0000"
    pdf = MIMEApplication(b"%PDF-1.4 mock pdf data", _subtype="pdf")
    pdf.add_header("Content-Disposition", "attachment", filename="stmt.pdf")
    msg.attach(pdf)

    mock_client = MagicMock()
    mock_client.select.return_value = ("OK", [b"1"])
    # Return same message ID in both folders to test deduplication across folders
    mock_client.search.return_value = ("OK", [b"101"])
    mock_client.fetch.return_value = ("OK", [(b"101 (RFC822 {1234}", msg.as_bytes())])

    results = fetch_and_process_statements(
        root_dir=tmp_path,
        days_back=30,
        folders=["INBOX", "Bank Statements"],
        client=mock_client,
    )

    # First folder downloads it; second folder skips because hash is already known
    assert len(results) == 2
    assert results[0].status == "downloaded"
    assert results[1].status == "skipped_hash"

    # Verify select was called for both folders. imaplib sends the mailbox name
    # as-is, so a name containing a space has to arrive already quoted.
    select_calls = [c[0][0] for c in mock_client.select.call_args_list]
    assert "INBOX" in select_calls
    assert '"Bank Statements"' in select_calls


def test_is_non_credit_card_email():
    from creditcard.email_fetcher import is_non_credit_card_email

    # Unwanted non-credit-card emails
    assert is_non_credit_card_email(
        "customernotification@icici.bank.in",
        "Transaction e-Statement for ICICI Bank Demat Account IN303028XXXXXX14",
    ) is True
    assert is_non_credit_card_email(
        "cbssbi.cas@alerts.sbi.bank.in",
        "E-account statement for your SBI account(s).",
    ) is True
    assert is_non_credit_card_email(
        "rewards@bank.com",
        "E-statement: You have 50 Points worth Rs.13 as on 30, June 2026",
    ) is True
    assert is_non_credit_card_email(
        "alerts@sbicard.com",
        "Transaction Alert from SBI Elite Master Card",
    ) is True

    # Genuine credit card emails
    assert is_non_credit_card_email(
        "statements@sbicard.com",
        "Your SBI Card ELITE Monthly Statement -Sep 2026",
    ) is False
    assert is_non_credit_card_email(
        "statements@sbicard.com",
        "Your BPCL SBI Card Monthly Statement -Sep 2026",
    ) is False
    assert is_non_credit_card_email(
        "alerts@hdfcbank.net",
        "Your HDFC Bank - Regalia HDFC Bank Credit Card Statement - Sep 2026",
    ) is False
    assert is_non_credit_card_email(
        "alerts@hdfcbank.net",
        "Your HDFC Bank - Millennia Infinity HDFC Bank Credit Card Statement - Sep 2026",
    ) is False
    assert is_non_credit_card_email(
        "credit_cards@icicibank.com",
        "ICICI Coral Bank Credit Card Statement for the period Jul 2026",
    ) is False


def test_export_fetch_results_csv_and_excel(tmp_path: Path):
    import csv

    import openpyxl

    csv_path = tmp_path / "fetched.csv"
    excel_path = tmp_path / "fetched.xlsx"

    results = [
        DownloadResult(
            email_date="2026-09-04",
            card_id="sbi_elite",
            status="downloaded",
            filename="E-Statement_092026.pdf",
            sender="statements@sbicard.com",
            subject="Your SBI Card ELITE Monthly Statement -Sep 2026",
            saved_path=str(tmp_path / "statements" / "sbi_elite" / "E-Statement_092026.pdf"),
            file_hash="dummyhash123",
            error="",
        )
    ]

    c_out, x_out = export_fetch_results(results, csv_path, excel_path)
    assert c_out == csv_path
    assert x_out == excel_path
    assert csv_path.exists()
    assert excel_path.exists()

    # Read CSV
    with open(csv_path, newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
        assert len(rows) == 1
        assert rows[0]["card_id"] == "sbi_elite"
        assert rows[0]["status"] == "downloaded"
        assert rows[0]["filename"] == "E-Statement_092026.pdf"

    # Read Excel
    wb = openpyxl.load_workbook(excel_path)
    assert "Fetched_Statements" in wb.sheetnames
    ws = wb["Fetched_Statements"]
    assert ws.cell(row=1, column=1).value == "email_date"
    assert ws.cell(row=2, column=2).value == "sbi_elite"
    assert ws.cell(row=2, column=3).value == "downloaded"
    assert ws.cell(row=2, column=4).value == "E-Statement_092026.pdf"



def _imap_cfg():

    return EmailConfig(host="imap.example.com", port=993,
                       username="me@example.com", password="app-password")


def test_rejected_login_is_an_auth_error(monkeypatch: pytest.MonkeyPatch):
    """A wrong app password won't fix itself, so the daily check reports it at
    once instead of waiting for a second failed run."""
    import imaplib

    from creditcard.email_fetcher import EmailAuthError, connect_imap

    class RejectingClient:
        def __init__(self, host, port):
            pass

        def login(self, user, password):
            raise imaplib.IMAP4.error("[AUTHENTICATIONFAILED] Invalid credentials")

        def logout(self):
            pass

    monkeypatch.setattr(imaplib, "IMAP4_SSL", RejectingClient)
    with pytest.raises(EmailAuthError) as info:
        connect_imap(_imap_cfg())
    assert "app-password" not in str(info.value)


def test_unreachable_server_is_a_plain_connection_error(monkeypatch: pytest.MonkeyPatch):
    import imaplib

    from creditcard.email_fetcher import EmailAuthError, connect_imap

    def unreachable(host, port):
        raise OSError("network is unreachable")

    monkeypatch.setattr(imaplib, "IMAP4_SSL", unreachable)
    with pytest.raises(ConnectionError) as info:
        connect_imap(_imap_cfg())
    assert not isinstance(info.value, EmailAuthError)
